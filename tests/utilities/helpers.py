from openpyxl import load_workbook
from selenium.webdriver.common.by import By


def read_data_from_excel(file_name, sheet):
    # Retrieves the data list from the manifest Excel sheet
    datalist = []
    workbook = load_workbook(filename=file_name)

    sheet = workbook.get_sheet_by_name(sheet)  # workbook[sheet]
    row_count = sheet.max_row
    column_count = sheet.max_column

    for i in range(2, row_count + 1):
        row = []
        for j in range(1, column_count + 1):
            row.append(sheet.cell(row=i, column=j).value)
            datalist.append(row)
    return datalist


def one_of_these_elements_is_visible(element1, element2):
    def wait_for_condition(driver):
        attribute = driver.find_element(By.ID, element1)
        if attribute.is_displayed():
            return attribute
        elif driver.find_element(By.ID, element2).is_displayed():
            return driver.find_element(By.ID, element2)

        else:
            return False

    return wait_for_condition


# Create a DummySMTP class that replaces smtplib.SMTP so that actual emails are not sent
smtp = None
inbox = []


class Message(object):
    def __init__(self, from_address, to_address, fullmessage):
        self.from_address = from_address
        self.to_address = to_address
        self.fullmessage = fullmessage


class DummySMTP(object):
    def __init__(self, address):
        self.address = address
        global smtp
        smtp = self

    def login(self, username, password):
        self.username = username
        self.password = password

    def sendmail(self, from_address, to_address, fullmessage):
        global inbox
        inbox.append(Message(from_address, to_address, fullmessage))
        return []

    def quit(self):
        self.has_quit = True
